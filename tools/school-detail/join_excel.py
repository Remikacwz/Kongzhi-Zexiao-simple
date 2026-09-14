# -*- coding: utf-8 -*-
"""方案A：把 Excel《录取数据表》的「拟录取人数(宝典)」并入 zxb_parsed.json 的「26实际录取」

匹配键（按优先级）：
  1. 学校名称 + 学院名称 + 专业代码   ← 同一代码跨学院时最准
  2. 学校名称 + 专业代码              ← 退一步，求和
专业代码从 Excel「专业名称」列的前 6 位提取。
"""
import json, os, re, sys, collections
import openpyxl

PARSED = os.environ.get('ZXB_PARSED_IN', '.qrcheck/zxb_parsed_fixed.json')
EXCEL  = os.environ.get('ZXB_EXCEL', '数据库/raw/27考研择校宝典_录取数据表_0815.xlsx')
OUT    = os.environ.get('ZXB_PARSED_OUT', '.qrcheck/zxb_joined.json')
CODE6  = re.compile(r'(\d[0-9A-Za-z]{5})')



def load_excel(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    merged = collections.defaultdict(int)      # (学校, 代码) -> 拟录取人数
    merged_c = collections.defaultdict(int)    # (学校, 学院, 代码) -> 拟录取人数
    per_school = collections.defaultdict(set)  # 学校 -> {代码}
    known_c = set()                            # (学校,学院,代码) 在 Excel 里出现过（含非数字值）
    known_s = set()                            # (学校,代码) 在 Excel 里出现过
    raw_rows = 0
    for r in rows:
        school = (r[0] or '').strip()
        college = (r[1] or '').strip()
        major  = (r[2] or '').strip()
        admit  = r[5]
        if not school or not major: continue
        m = CODE6.search(major)
        if not m: continue
        code = m.group(1)
        per_school[school].add(code)
        known_s.add((school, code))
        if college: known_c.add((school, col_key(college), code))
        if isinstance(admit, (int, float)):
            merged[(school, code)] += int(admit)
            if college:
                merged_c[(school, col_key(college), code)] += int(admit)
            raw_rows += 1
    return merged, merged_c, per_school, known_c, known_s, len(rows), raw_rows


def col_key(name):
    """学院名归一化：去掉前导数字代码、括号说明、尾部标点"""
    c = re.sub(r'^[（(]?\d{2,4}[）)]?[\s、]*', '', (name or '').strip())
    c = re.sub(r'[（(][^（()）]*[）)]', '', c)
    return re.sub(r'[：:\s、，,]+$', '', c).strip()


def main():
    parsed = json.loads(open(PARSED, encoding='utf-8').read())
    merged, merged_c, per_school, known_c, known_s, nrows, nraw = load_excel(EXCEL)
    hit = miss = hit_c = 0
    misslist = []
    for s in parsed:
        for p in s['专业']:
            key  = (s['学校'], p['代码'])
            keyc = (s['学校'], col_key(p.get('学院', '')), p['代码'])
            val = None
            if keyc in merged_c:                              # 1) 学院级精确匹配
                val = merged_c[keyc]; hit_c += 1
            else:                                             # 1b) 学院名子串匹配
                sub = [v for (sc, ck, cd), v in merged_c.items()
                       if sc == s['学校'] and cd == p['代码'] and ck and keyc[1]
                       and (ck in keyc[1] or keyc[1] in ck)]
                if len(sub) == 1:
                    val = sub[0]; hit_c += 1
                elif key in merged and keyc not in known_c \
                        and not any(sc == s['学校'] and cd == p['代码'] and ck and keyc[1]
                                    and (ck in keyc[1] or keyc[1] in ck)
                                    for (sc, ck, cd) in known_c):     # 2) 学校级求和
                    val = merged[key]
                elif p['实际录取'] is None and key not in known_s:   # 3) 同校前缀代码
                    alt = [v for (sc, cd), v in merged.items()
                           if sc == s['学校'] and cd.startswith(p['代码'][:4])]
                    if alt: val = sum(alt)
            if val is not None:
                p['实际录取'] = val; hit += 1
            elif p['实际录取'] is None:
                miss += 1; misslist.append((s['学校'], p['代码'], p['名称'], p['学院']))
    open(OUT, 'w', encoding='utf-8').write(json.dumps(parsed, ensure_ascii=False, indent=1))
    print(f'Excel 数据行 {nrows}  有拟录取人数 {nraw}')
    print(f'并入成功 {hit}（学院级精确匹配 {hit_c}）   未能匹配 {miss}   ({hit/(hit+miss)*100:.1f}%)')
    print(f'输出: {OUT}')
    print()
    print('未匹配（全部）：')


if __name__ == '__main__':
    main()
