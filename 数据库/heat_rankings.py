# -*- coding: utf-8 -*-
"""热度榜 Excel 解析、院校匹配与榜单持久化。"""
from __future__ import annotations

import base64
import binascii
import io
import pathlib
import re
import sqlite3
import unicodedata


BASE_DIR = pathlib.Path(__file__).resolve().parent.parent
MAX_EXCEL_BYTES = 12 * 1024 * 1024
SHEET_SCOPES = (
    ('总热度榜', 'all', '全部'),
    ('985热度榜', '985', '985'),
    ('211热度榜', '211', '211'),
    ('普通院校热度榜', 'double_non', '双非'),
)
PUBLIC_SCOPES = {'all', '985', '211', 'double_non'}


def ensure_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS heat_ranking_batches (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          period TEXT NOT NULL UNIQUE,
          source_filename TEXT NOT NULL,
          imported_by INTEGER,
          created_at TEXT NOT NULL,
          published_at TEXT NOT NULL,
          FOREIGN KEY (imported_by) REFERENCES admin_users(id) ON DELETE SET NULL
        );
        CREATE TABLE IF NOT EXISTS heat_ranking_rows (
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          batch_id INTEGER NOT NULL,
          period TEXT NOT NULL,
          scope TEXT NOT NULL,
          rank INTEGER NOT NULL,
          source_school_name TEXT NOT NULL,
          school_name TEXT NOT NULL,
          school_id INTEGER,
          heat REAL NOT NULL,
          total_rank INTEGER,
          tier TEXT NOT NULL,
          source_tier TEXT NOT NULL DEFAULT '',
          created_at TEXT NOT NULL,
          updated_at TEXT NOT NULL,
          UNIQUE(period, scope, rank),
          UNIQUE(period, scope, source_school_name),
          FOREIGN KEY (batch_id) REFERENCES heat_ranking_batches(id) ON DELETE CASCADE,
          FOREIGN KEY (school_id) REFERENCES schools(id) ON DELETE SET NULL
        );
        CREATE INDEX IF NOT EXISTS idx_heat_ranking_public
          ON heat_ranking_rows(period, scope, rank);
        """
    )


def normalize_period(value: str) -> str:
    digits = re.sub(r'\D', '', str(value or ''))
    if len(digits) != 6:
        raise ValueError('请选择榜单所属年月')
    year, month = int(digits[:4]), int(digits[4:])
    if year < 2020 or year > 2100 or month < 1 or month > 12:
        raise ValueError('榜单年月格式不正确')
    return f'{year:04d}{month:02d}'


def _name_key(value: str) -> str:
    text = unicodedata.normalize('NFKC', str(value or '')).strip()
    text = text.replace('（', '(').replace('）', ')')
    return re.sub(r'[\s·•・—_\-()（）]', '', text).lower()


def _school_catalog(conn: sqlite3.Connection) -> list[dict]:
    merged: dict[str, dict] = {}
    try:
        rows = conn.execute('SELECT id,name,province,tier FROM schools ORDER BY name').fetchall()
    except sqlite3.OperationalError:
        rows = []
    for row in rows:
        item = dict(row)
        name = str(item.get('name') or '').strip()
        if name:
            merged[name] = {
                'name': name,
                'school_id': int(item['id']),
                'province': str(item.get('province') or ''),
                'tier': str(item.get('tier') or ''),
            }

    detail_dirs = (
        BASE_DIR / 'school_detail',
        BASE_DIR / '就业相关' / '院校就业去向' / 'schools',
    )
    for detail_dir in detail_dirs:
        if not detail_dir.is_dir():
            continue
        for path in detail_dir.glob('*.html'):
            if path.stem.lower() == 'index':
                continue
            name = path.stem.strip()
            if name and name not in merged:
                merged[name] = {'name': name, 'school_id': None, 'province': '', 'tier': ''}
    return sorted(merged.values(), key=lambda item: item['name'])


def _match_school(source_name: str, catalog: list[dict]) -> dict:
    source_key = _name_key(source_name)
    keyed: dict[str, list[dict]] = {}
    for school in catalog:
        keyed.setdefault(_name_key(school['name']), []).append(school)

    exact = keyed.get(source_key, [])
    if len(exact) > 1:
        return _match_payload('ambiguous', None, exact)

    # 榜单常省略校区，例如“华北电力大学”或“中国石油大学”。即使存在
    # 一个通用同名页面，也必须先检查“（北京）/（保定）”等具体校区。
    descendants = []
    for school in catalog:
        candidate_key = _name_key(school['name'])
        if candidate_key != source_key and candidate_key.startswith(source_key):
            descendants.append(school)
    descendants = list({item['name']: item for item in descendants}.values())
    if len(descendants) > 1:
        return _match_payload('ambiguous', None, descendants)
    if len(descendants) == 1 and exact:
        return _match_payload('ambiguous', None, [exact[0], descendants[0]])
    if len(descendants) == 1:
        return _match_payload('matched', descendants[0], descendants)
    if len(exact) == 1:
        return _match_payload('matched', exact[0], exact)

    # 兼容来源名称比目录名称更长的少量情况；仍然只在候选唯一时自动确认。
    parents = []
    for school in catalog:
        candidate_key = _name_key(school['name'])
        if candidate_key and source_key.startswith(candidate_key):
            parents.append(school)
    parents = list({item['name']: item for item in parents}.values())
    if len(parents) == 1:
        return _match_payload('matched', parents[0], parents)
    if len(parents) > 1:
        return _match_payload('ambiguous', None, parents)
    return _match_payload('unmatched', None, [])


def _match_payload(status: str, selected: dict | None, candidates: list[dict]) -> dict:
    return {
        'status': status,
        'school_name': selected['name'] if selected else '',
        'school_id': selected.get('school_id') if selected else None,
        'candidates': [
            {'name': item['name'], 'school_id': item.get('school_id')}
            for item in candidates
        ],
    }


def _number(value, label: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        raise ValueError(f'{label}必须是数字') from None


def _integer(value, label: str) -> int:
    number = _number(value, label)
    if int(number) != number:
        raise ValueError(f'{label}必须是整数')
    return int(number)


def _tier_for(scope: str, source_tier: str) -> str:
    if scope == '985' or '985' in source_tier:
        return '985'
    if scope == '211' or '211' in source_tier:
        return '211'
    return '双非'


def _read_scope_sheet(sheet, sheet_name: str, scope: str, catalog: list[dict]) -> list[dict]:
    header = [str(sheet.cell(1, col).value or '').strip() for col in range(1, 5)]
    if len(header) < 3 or '排名' not in header[0] or header[1] != '招生单位' or header[2] != '热度值':
        raise ValueError(f'工作表“{sheet_name}”表头不符合要求，应包含排名、招生单位、热度值')

    source_rows = []
    for row_index, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        rank_value, school_value, heat_value = values[:3]
        if rank_value in (None, '') and school_value in (None, ''):
            continue
        school_name = str(school_value or '').strip()
        if not school_name:
            raise ValueError(f'工作表“{sheet_name}”第 {row_index} 行缺少招生单位')
        source_rank = _integer(rank_value, f'工作表“{sheet_name}”第 {row_index} 行排名')
        heat = _number(heat_value, f'工作表“{sheet_name}”第 {row_index} 行热度值')
        source_rows.append((source_rank, row_index, school_name, heat, values[3] if len(values) > 3 else None))

    source_rows.sort(key=lambda item: (item[0], item[1]))
    selected = source_rows[:20]
    if len(selected) < 20:
        raise ValueError(f'工作表“{sheet_name}”有效数据不足 20 行')

    items = []
    for display_rank, (_source_rank, row_index, school_name, heat, fourth) in enumerate(selected, start=1):
        source_tier = str(fourth or '').strip() if scope == 'all' else ''
        total_rank = display_rank if scope == 'all' else (_integer(fourth, f'工作表“{sheet_name}”第 {row_index} 行总排名') if fourth not in (None, '') else None)
        items.append({
            'rank': display_rank,
            'source_rank': _source_rank,
            'source_school_name': school_name,
            'heat': round(heat, 4),
            'total_rank': total_rank,
            'tier': _tier_for(scope, source_tier),
            'source_tier': source_tier,
            'match': _match_school(school_name, catalog),
        })
    return items


def preview_excel(conn: sqlite3.Connection, body: dict) -> dict:
    period = normalize_period(body.get('period', ''))
    filename = pathlib.Path(str(body.get('filename') or '热度榜单.xlsx')).name
    if not filename.lower().endswith('.xlsx'):
        raise ValueError('仅支持 .xlsx 文件')
    encoded = str(body.get('base64') or '')
    if not encoded:
        raise ValueError('请选择热度榜 Excel 文件')
    try:
        raw = base64.b64decode(encoded, validate=True)
    except (ValueError, binascii.Error):
        raise ValueError('Excel 文件内容无效') from None
    if not raw or len(raw) > MAX_EXCEL_BYTES:
        raise ValueError('Excel 文件不能为空且不能超过 12MB')

    try:
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f'无法读取 Excel：{exc}') from None

    missing = [name for name, _scope, _label in SHEET_SCOPES if name not in workbook.sheetnames]
    if missing:
        raise ValueError('缺少工作表：' + '、'.join(missing))

    catalog = _school_catalog(conn)
    groups = []
    matched = ambiguous = unmatched = 0
    try:
        for sheet_name, scope, label in SHEET_SCOPES:
            items = _read_scope_sheet(workbook[sheet_name], sheet_name, scope, catalog)
            for item in items:
                status = item['match']['status']
                matched += status == 'matched'
                ambiguous += status == 'ambiguous'
                unmatched += status == 'unmatched'
            groups.append({
                'scope': scope,
                'label': label,
                'sheet_name': sheet_name,
                'items': items,
            })
    finally:
        workbook.close()

    return {
        'period': period,
        'year': int(period[:4]),
        'month': int(period[4:]),
        'filename': filename,
        'groups': groups,
        'catalog': catalog,
        'summary': {
            'rows': sum(len(group['items']) for group in groups),
            'matched': matched,
            'ambiguous': ambiguous,
            'unmatched': unmatched,
        },
    }


def publish_preview(conn: sqlite3.Connection, body: dict, user_id: int, now: str) -> dict:
    period = normalize_period(body.get('period', ''))
    filename = pathlib.Path(str(body.get('filename') or '热度榜单.xlsx')).name
    groups = body.get('groups')
    if not isinstance(groups, list):
        raise ValueError('缺少榜单预览数据，请重新解析 Excel')

    catalog = _school_catalog(conn)
    catalog_by_name = {item['name']: item for item in catalog}
    groups_by_scope = {str(group.get('scope')): group for group in groups if isinstance(group, dict)}
    if set(groups_by_scope) != PUBLIC_SCOPES:
        raise ValueError('榜单必须同时包含全部、985、211、双非四个分组')

    normalized_groups = []
    for scope in ('all', '985', '211', 'double_non'):
        items = groups_by_scope[scope].get('items')
        if not isinstance(items, list) or len(items) != 20:
            raise ValueError(f'{scope} 榜单必须包含前 20 名')
        ranks = set()
        source_names = set()
        normalized_items = []
        for item in items:
            rank = _integer(item.get('rank'), f'{scope} 排名')
            if rank < 1 or rank > 20 or rank in ranks:
                raise ValueError(f'{scope} 榜单排名必须是互不重复的 1—20')
            ranks.add(rank)
            source_name = str(item.get('source_school_name') or '').strip()
            school_name = str(item.get('school_name') or '').strip()
            if source_name in source_names:
                raise ValueError(f'{scope} 榜单存在重复院校“{source_name}”')
            source_names.add(source_name)
            if not source_name or not school_name or school_name not in catalog_by_name:
                raise ValueError(f'{scope} 第 {rank} 名“{source_name or "未命名院校"}”尚未完成院校匹配')
            school = catalog_by_name[school_name]
            source_tier = str(item.get('source_tier') or '').strip()
            normalized_items.append({
                'rank': rank,
                'source_school_name': source_name,
                'school_name': school_name,
                'school_id': school.get('school_id'),
                'heat': _number(item.get('heat'), f'{scope} 第 {rank} 名热度值'),
                'total_rank': _integer(item.get('total_rank'), f'{scope} 第 {rank} 名总排名') if item.get('total_rank') not in (None, '') else None,
                'tier': _tier_for(scope, source_tier),
                'source_tier': source_tier,
            })
        normalized_groups.append((scope, sorted(normalized_items, key=lambda item: item['rank'])))

    old = conn.execute('SELECT id FROM heat_ranking_batches WHERE period=?', (period,)).fetchone()
    if old:
        conn.execute('DELETE FROM heat_ranking_batches WHERE id=?', (old['id'],))
    cursor = conn.execute(
        'INSERT INTO heat_ranking_batches(period,source_filename,imported_by,created_at,published_at) VALUES(?,?,?,?,?)',
        (period, filename, user_id, now, now),
    )
    batch_id = int(cursor.lastrowid)
    for scope, items in normalized_groups:
        conn.executemany(
            """INSERT INTO heat_ranking_rows(
                 batch_id,period,scope,rank,source_school_name,school_name,school_id,
                 heat,total_rank,tier,source_tier,created_at,updated_at
               ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            [(
                batch_id, period, scope, item['rank'], item['source_school_name'], item['school_name'],
                item['school_id'], item['heat'], item['total_rank'], item['tier'], item['source_tier'], now, now,
            ) for item in items],
        )
    return {'batch_id': batch_id, 'period': period, 'rows': 80, 'replaced': bool(old)}


def public_rankings(conn: sqlite3.Connection, period: str = '', scope: str = 'all', limit=20) -> dict:
    ensure_schema(conn)
    requested_scope = str(scope or 'all')
    if requested_scope == '双非':
        requested_scope = 'double_non'
    if requested_scope not in PUBLIC_SCOPES:
        raise ValueError('不支持的榜单层级')
    try:
        requested_limit = max(1, min(20, int(limit or 20)))
    except (TypeError, ValueError):
        requested_limit = 20

    periods = [row['period'] for row in conn.execute(
        'SELECT period FROM heat_ranking_batches ORDER BY period DESC'
    ).fetchall()]
    selected_period = normalize_period(period) if period else (periods[0] if periods else '')
    if not selected_period:
        return {'period': '', 'scope': requested_scope, 'items': [], 'available_periods': []}

    rows = conn.execute(
        """SELECT rank,source_school_name,school_name,school_id,heat,total_rank,tier
           FROM heat_ranking_rows
           WHERE period=? AND scope=? ORDER BY rank LIMIT ?""",
        (selected_period, requested_scope, requested_limit),
    ).fetchall()
    return {
        'period': selected_period,
        'year': int(selected_period[:4]),
        'month': int(selected_period[4:]),
        'scope': requested_scope,
        'items': [dict(row) for row in rows],
        'available_periods': periods,
    }
